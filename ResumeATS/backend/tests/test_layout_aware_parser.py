from __future__ import annotations

import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch
from zipfile import ZipFile

from docx import Document

from app.parser_pipeline import pipeline
from app.resume_io import read_resume_text


class LayoutAwareResumeParserTests(unittest.TestCase):
    def _word_resume_bytes(self) -> bytes:
        doc = Document()
        doc.add_paragraph("EXPERIENCE SUMMARY:")
        doc.add_paragraph("Senior database professional supporting secure systems.")
        doc.add_paragraph("EDUCATION | CLEARANCES:")
        p = doc.add_paragraph()
        p.add_run("Example University")
        p.add_run().add_break()
        p.add_run("Bachelor of Science: Information Systems")
        doc.add_paragraph("TECHNICAL SKILL SETS:")
        doc.add_paragraph("RDBMS:\tOracle, SQL Server, PostgreSQL")
        doc.add_paragraph("PROFESSIONAL EXPERIENCE:")
        p = doc.add_paragraph()
        p.add_run("Example Consulting Group, LLC\t\t06/2015 - current")
        p.add_run().add_break()
        p.add_run("Title: Database Consultant")
        doc.add_paragraph("Provided database architecture and production support.")
        doc.add_paragraph("Federal Security Operations; Arlington, VA\t02/2013 - 06/2015")
        doc.add_paragraph("Title: Database Team Lead | SME Contract: Example Systems")
        doc.add_paragraph("Environment: Classified systems and secure databases")
        doc.add_paragraph("PROFESSIONAL EXPERIENCE (CONT)…")
        doc.add_paragraph("U.S. Government Program; Washington, DC\t06/2009 - 02/2013")
        doc.add_paragraph("Title: Oracle DBA | Team Lead & SharePoint Administrator Contract: Example Digital, LLC")
        doc.add_paragraph("Managed database operations across secure enterprise systems.")
        doc.add_paragraph("REFERENCES: Furnished upon request")

        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def test_ingestion_preserves_manual_breaks_and_tabs(self):
        raw = self._word_resume_bytes()
        file = SimpleNamespace(filename="resume.docx")
        text = read_resume_text(file, raw)

        self.assertIn("Example Consulting Group, LLC | 06/2015 - current", text)
        self.assertIn("\nTitle: Database Consultant\n", text)
        self.assertIn("Example University\nBachelor of Science", text)

    @patch("app.parser_pipeline.classifier.organization_match", return_value={"valid": False, "confidence": 0.0})
    @patch("app.parser_pipeline.classifier.job_title_match", return_value={"valid": False, "confidence": 0.0})
    def test_non_ats_word_layout_extracts_sections_and_jobs(self, _title_match, _org_match):
        raw = self._word_resume_bytes()
        file = SimpleNamespace(filename="resume.docx")
        text = read_resume_text(file, raw)
        result = pipeline.parse(text)

        self.assertIn("database professional", result.executive_summary.lower())
        self.assertIn("Oracle", result.skills)
        self.assertEqual(len(result.education), 1)
        self.assertIn("Example University", result.education[0])
        self.assertGreaterEqual(len(result.experience), 3)
        self.assertEqual(result.experience[0].job, "Database Consultant")
        self.assertEqual(result.experience[0].company, "Example Consulting Group, LLC")
        self.assertEqual(result.experience[0].date_range, "06/2015 - current")
        self.assertEqual(result.experience[1].job, "Database Team Lead")
        self.assertIn("Federal Security Operations", result.experience[1].company)
        self.assertIn("Oracle DBA", result.experience[2].job)
        self.assertNotIn("Furnished upon request", " ".join(result.experience[-1].descriptions))
        self.assertEqual(result.target_position_title, "Database Consultant")

    def test_readable_resume_formats(self):
        html = b"<html><body><h1>Jane Example</h1><p>Database Engineer</p><script>ignore this</script></body></html>"
        self.assertIn("Jane Example", read_resume_text(SimpleNamespace(filename="resume.html"), html))
        self.assertNotIn("ignore this", read_resume_text(SimpleNamespace(filename="resume.html"), html))

        json_resume = b'{"name": "Jane Example", "skills": ["SQL", "Python"]}'
        json_text = read_resume_text(SimpleNamespace(filename="resume.json"), json_resume)
        self.assertIn("name: Jane Example", json_text)
        self.assertIn("skills: SQL", json_text)

        xml_resume = b"<resume><name>Jane Example</name><summary>Database Engineer</summary></resume>"
        xml_text = read_resume_text(SimpleNamespace(filename="resume.xml"), xml_resume)
        self.assertIn("Jane Example", xml_text)
        self.assertIn("Database Engineer", xml_text)

        odt_content = b'''<?xml version="1.0" encoding="UTF-8"?>
        <office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
          xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
          <office:body><office:text><text:h>Jane Example</text:h>
          <text:p>Database<text:s/>Engineer</text:p></office:text></office:body>
        </office:document-content>'''
        odt_buffer = BytesIO()
        with ZipFile(odt_buffer, "w") as archive:
            archive.writestr("content.xml", odt_content)
        odt_text = read_resume_text(SimpleNamespace(filename="resume.odt"), odt_buffer.getvalue())
        self.assertIn("Jane Example", odt_text)
        self.assertIn("Database Engineer", odt_text)

        archive_buffer = BytesIO()
        with ZipFile(archive_buffer, "w") as archive:
            archive.writestr("resume.xml", xml_resume)
        archive_text = read_resume_text(SimpleNamespace(filename="resume.zip"), archive_buffer.getvalue())
        self.assertIn("Jane Example", archive_text)

        pages_buffer = BytesIO()
        with ZipFile(pages_buffer, "w") as archive:
            archive.writestr("Index.xml", xml_resume)
        pages_text = read_resume_text(SimpleNamespace(filename="resume.pages"), pages_buffer.getvalue())
        self.assertIn("Database Engineer", pages_text)

    @patch("app.resume_io.shutil.which", side_effect=lambda name: "/usr/bin/antiword" if name == "antiword" else None)
    @patch("app.resume_io.subprocess.run")
    def test_legacy_doc_uses_installed_converter(self, run, _which):
        run.return_value = Mock(stdout=b"Jane Example\nDatabase Engineer\n")
        text = read_resume_text(SimpleNamespace(filename="resume.doc"), b"legacy word bytes")
        self.assertIn("Jane Example", text)
        run.assert_called_once()


if __name__ == "__main__":
    unittest.main()
